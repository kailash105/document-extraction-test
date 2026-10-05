import io
import zipfile
import pytest
from docx import Document
from openpyxl import Workbook
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from app.config import ROOT
from app.readers import DocumentError, prepare_document


@pytest.mark.parametrize('name,data,expected', [
    ('a.txt', b'Rent is 12000', '12000'),
    ('a.md', b'# Agreement\nParty A', 'Party A'),
    ('a.csv', b'Party,Value\nAlice,12000', 'Alice | 12000'),
    ('a.tsv', b'Party\tValue\nAlice\t12000', 'Alice | 12000'),
    ('a.json', b'{"party":"Alice"}', 'Alice'),
    ('a.html', b'<script>SECRET</script><p>Alice</p>', 'Alice'),
    ('a.xml', b'<agreement><party>Alice</party></agreement>', 'Alice'),
    ('a.rtf', b'{\\rtf1\\ansi Alice pays 12000}', 'Alice'),
    ('a.eml', b'Subject: Agreement\nContent-Type: text/plain\n\nAlice pays 12000', 'Alice'),
    ('a.txt', 'Rent €12000'.encode('utf-16'), '€12000'),
])
def test_text_formats(name, data, expected):
    doc = prepare_document(name, data)
    assert expected in doc.text
    assert 'SECRET' not in doc.text


def test_docx_tables_headers_and_embedded_image():
    document = Document()
    document.add_paragraph('Rental agreement')
    document.add_table(rows=1, cols=2).rows[0].cells[1].text = 'Alice'
    document.sections[0].header.paragraphs[0].text = 'Important header'
    image = io.BytesIO()
    Image.new('RGB', (50, 50), 'white').save(image, 'PNG')
    document.add_picture(io.BytesIO(image.getvalue()))
    stream = io.BytesIO(); document.save(stream)
    doc = prepare_document('x.docx', stream.getvalue())
    assert 'Alice' in doc.text and 'Important header' in doc.text
    assert len(doc.images) == 1


def test_xlsx_multiple_sheets():
    book = Workbook(); book.active.append(['Party', 'Alice'])
    book.create_sheet('Other').append(['Rent', 12000])
    stream = io.BytesIO(); book.save(stream)
    doc = prepare_document('x.xlsx', stream.getvalue())
    assert 'Alice' in doc.text and '12000' in doc.text and doc.units == 2


def test_pptx():
    deck = Presentation(); slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1)).text = 'Alice pays 12000'
    stream = io.BytesIO(); deck.save(stream)
    assert 'Alice pays 12000' in prepare_document('x.pptx', stream.getvalue()).text


def test_odt():
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, 'w') as archive:
        archive.writestr('content.xml', '<document><paragraph>Alice pays 12000</paragraph></document>')
    assert 'Alice' in prepare_document('x.odt', stream.getvalue()).text


@pytest.mark.parametrize('ext,format', [('png','PNG'), ('jpg','JPEG'), ('webp','WEBP'), ('bmp','BMP'), ('gif','GIF'), ('tiff','TIFF')])
def test_image_formats(ext, format):
    stream = io.BytesIO(); Image.new('RGB', (100, 3000), 'white').save(stream, format)
    doc = prepare_document('scan.' + ext, stream.getvalue())
    assert len(doc.images) == 2  # Tall scan preserves details through tiling.


def test_pdf_all_pages_including_scans():
    stream = io.BytesIO()
    Image.new('RGB', (200, 300), 'white').save(stream, 'PDF', save_all=True, append_images=[Image.new('RGB', (200, 300), 'black')])
    doc = prepare_document('scan.pdf', stream.getvalue())
    assert doc.units == 2 and len(doc.images) == 2


def test_multipage_tiff():
    stream = io.BytesIO()
    Image.new('RGB', (100, 100), 'white').save(stream, 'TIFF', save_all=True, append_images=[Image.new('RGB', (100, 100), 'black')])
    doc = prepare_document('scan.tif', stream.getvalue())
    assert doc.units == 2 and len(doc.images) == 2


@pytest.mark.parametrize('name,data', [('bad.docx', b'not a zip'), ('bad.png', b'not an image'), ('old.doc', b'old format'), ('empty.txt', b''), ('binary.txt', b'\x00\x01')])
def test_invalid_files_fail_clearly(name, data):
    with pytest.raises(DocumentError):
        prepare_document(name, data)


def test_xml_entities_are_rejected():
    with pytest.raises(DocumentError):
        prepare_document('unsafe.xml', b'<!DOCTYPE a [<!ENTITY x SYSTEM "file:///etc/passwd">]><a>&x;</a>')


def test_size_and_text_limits(monkeypatch):
    monkeypatch.setattr('app.readers.MAX_FILE_BYTES', 20)
    with pytest.raises(DocumentError, match='20 MB'):
        prepare_document('a.txt', b'A' * 21)
    monkeypatch.setattr('app.readers.MAX_TEXT_CHARS', 5)
    with pytest.raises(DocumentError, match='characters'):
        prepare_document('a.txt', b'A' * 10)


@pytest.mark.parametrize('path', sorted((ROOT / 'data').glob('*/*')))
def test_every_supplied_document_is_readable(path):
    doc = prepare_document(path.name, path.read_bytes())
    assert doc.text or doc.images
