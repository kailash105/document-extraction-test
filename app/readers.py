"""Format decoding only. No rules here select or extract agreement fields."""
import base64
import csv
import io
import json
import threading
import warnings
import zipfile
from dataclasses import dataclass, field
from email import policy
from email.parser import BytesParser
from pathlib import Path

import pypdfium2 as pdfium
from bs4 import BeautifulSoup
from defusedxml import ElementTree
from docx import Document
from docx.table import Table
from openpyxl import load_workbook
from PIL import Image, ImageOps
from pptx import Presentation
from striprtf.striprtf import rtf_to_text

from .config import MAX_ARCHIVE_BYTES, MAX_FILE_BYTES, MAX_IMAGES, MAX_PAGES, MAX_TEXT_CHARS

FORMATS = {
    'Documents': ['.pdf', '.docx', '.odt', '.rtf'],
    'Images': ['.png', '.jpg', '.jpeg', '.webp', '.tif', '.tiff', '.bmp', '.gif'],
    'Text & data': ['.txt', '.md', '.csv', '.tsv', '.json', '.html', '.htm', '.xml', '.eml'],
    'Office': ['.xlsx', '.pptx'],
}
EXTENSIONS = {ext for values in FORMATS.values() for ext in values}
PDF_LOCK = threading.Lock()  # PDFium is not thread safe.


class DocumentError(ValueError):
    pass


@dataclass
class PreparedDocument:
    name: str
    format: str
    text: str = ''
    images: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    units: int = 1

    def add_text(self, text: str):
        self.text += text + '\n'
        if len(self.text) > MAX_TEXT_CHARS:
            raise DocumentError(f'Document exceeds {MAX_TEXT_CHARS:,} text characters. Split it into smaller documents.')

    def add_image(self, im: Image.Image):
        # Slice tall scans instead of shrinking multiple pages into one unreadable image.
        im = ImageOps.exif_transpose(im).convert('RGB')
        if im.width > 1600:
            im = im.resize((1600, max(1, round(im.height * 1600 / im.width))))
        top = 0
        while top < im.height:
            if len(self.images) >= MAX_IMAGES:
                raise DocumentError(f'Document needs more than {MAX_IMAGES} image sections. Split it into smaller documents.')
            bottom = min(top + 2000, im.height)
            section = im.crop((0, top, im.width, bottom))
            out = io.BytesIO()
            section.save(out, format='JPEG', quality=88)
            self.images.append('data:image/jpeg;base64,' + base64.b64encode(out.getvalue()).decode())
            if bottom == im.height:
                break
            top = bottom - 100


def _decode(data: bytes) -> str:
    for encoding in ('utf-8-sig', 'utf-16' if data.startswith((b'\xff\xfe', b'\xfe\xff')) else 'utf-8', 'cp1252'):
        try:
            result = data.decode(encoding)
            if '\x00' in result:
                raise DocumentError('This file contains binary data, not readable text.')
            return result
        except UnicodeError:
            continue
    raise DocumentError('Cannot decode the text. Save the document using UTF-8.')


def _html(text: str) -> str:
    soup = BeautifulSoup(text, 'html.parser')
    for element in soup(['script', 'style', 'noscript']):
        element.decompose()
    return soup.get_text(separator='\n', strip=True)


def _check_archive(data: bytes):
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        members = archive.infolist()
        if len(members) > 5000 or sum(m.file_size for m in members) > MAX_ARCHIVE_BYTES:
            raise DocumentError('The expanded document is too large to read safely.')


def _archive_images(doc: PreparedDocument, data: bytes, prefix: str):
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        for name in sorted(archive.namelist()):
            if not name.startswith(prefix) or name.endswith('/'):
                continue
            if Path(name).suffix.lower() not in FORMATS['Images']:
                doc.warnings.append(f'Embedded media {Path(name).name} could not be read. Export to PDF to include it.')
                continue
            with Image.open(io.BytesIO(archive.read(name))) as im:
                doc.add_image(im)


def _read_pdf(doc: PreparedDocument, data: bytes):
    with PDF_LOCK:
        pdf = pdfium.PdfDocument(data)
        try:
            doc.units = len(pdf)
            if len(pdf) > MAX_PAGES:
                raise DocumentError(f'PDFs are limited to {MAX_PAGES} pages. Split this file first.')
            for index in range(len(pdf)):
                page = pdf[index]
                try:
                    textpage = page.get_textpage()
                    try:
                        doc.add_text(f'[Page {index + 1}]\n' + textpage.get_text_range())
                    finally:
                        textpage.close()
                    scale = min(2, 1600 / max(page.get_width(), 1))
                    if page.get_height() * scale > 16000:
                        raise DocumentError('PDF page is too tall. Split the page into smaller pages.')
                    bitmap = page.render(scale=scale)
                    try:
                        doc.add_image(bitmap.to_pil())
                    finally:
                        bitmap.close()
                finally:
                    page.close()
        finally:
            pdf.close()


def _read_docx(doc: PreparedDocument, data: bytes):
    document = Document(io.BytesIO(data))
    for block in document.iter_inner_content():
        if isinstance(block, Table):
            for row in block.rows:
                doc.add_text(' | '.join(cell.text for cell in row.cells))
        else:
            doc.add_text(block.text)
    for section in document.sections:
        for part in (section.header, section.footer):
            for paragraph in part.paragraphs:
                doc.add_text(paragraph.text)
    _archive_images(doc, data, 'word/media/')
    # Text boxes and footnotes are not exposed by python-docx's paragraph API.
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
        for name in archive.namelist():
            if name in ('word/footnotes.xml', 'word/endnotes.xml', 'word/document.xml'):
                root = ElementTree.fromstring(archive.read(name))
                nodes = root.findall('.//w:txbxContent', ns) if name.endswith('document.xml') else [root]
                for node in nodes:
                    doc.add_text(' '.join(t.text or '' for t in node.findall('.//w:t', ns)))


def _read_xlsx(doc: PreparedDocument, data: bytes):
    book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        doc.units = len(book.worksheets)
        count = 0
        for sheet in book:
            doc.add_text(f'[Sheet: {sheet.title}]')
            for row in sheet.iter_rows(values_only=True):
                count += len(row)
                if count > 100_000:
                    raise DocumentError('Workbook exceeds 100,000 cells. Export the relevant sheet as CSV.')
                if any(v is not None for v in row):
                    doc.add_text(' | '.join('' if v is None else str(v) for v in row))
    finally:
        book.close()
    _archive_images(doc, data, 'xl/media/')
    doc.warnings.append('Spreadsheet formulas use saved values. Charts and formatting are not rendered; export to PDF if they contain relevant information.')


def _read_pptx(doc: PreparedDocument, data: bytes):
    deck = Presentation(io.BytesIO(data))
    doc.units = len(deck.slides)
    if doc.units > MAX_PAGES:
        raise DocumentError(f'Presentations are limited to {MAX_PAGES} slides.')

    def read_shapes(shapes):
        for shape in shapes:
            if hasattr(shape, 'shapes'):
                read_shapes(shape.shapes)
            if shape.has_text_frame:
                doc.add_text(shape.text)
            if shape.has_table:
                for row in shape.table.rows:
                    doc.add_text(' | '.join(cell.text for cell in row.cells))

    for i, slide in enumerate(deck.slides, 1):
        doc.add_text(f'[Slide {i}]')
        read_shapes(slide.shapes)
        if slide.has_notes_slide:
            doc.add_text(slide.notes_slide.notes_text_frame.text)
    _archive_images(doc, data, 'ppt/media/')
    doc.warnings.append('Slide text, notes and embedded images are included. Charts and drawn shapes are not rendered; use PDF for full visual fidelity.')


def prepare_document(name: str, data: bytes) -> PreparedDocument:
    suffix = Path(name).suffix.lower()
    if suffix not in EXTENSIONS:
        raise DocumentError(f'Unsupported file type {suffix or "(no extension)"}. Convert legacy Office files to DOCX, XLSX, PPTX or PDF; other formats to PDF or text.')
    if not data:
        raise DocumentError('The file is empty.')
    if len(data) > MAX_FILE_BYTES:
        raise DocumentError('Files must be 20 MB or smaller.')
    doc = PreparedDocument(name=Path(name).name, format=suffix[1:])
    try:
        with warnings.catch_warnings():
            warnings.simplefilter('error', Image.DecompressionBombWarning)
            if suffix in {'.docx', '.xlsx', '.pptx', '.odt'}:
                _check_archive(data)
            if suffix == '.pdf':
                _read_pdf(doc, data)
            elif suffix == '.docx':
                _read_docx(doc, data)
            elif suffix in FORMATS['Images']:
                with Image.open(io.BytesIO(data)) as im:
                    doc.units = getattr(im, 'n_frames', 1)
                    if doc.units > MAX_PAGES:
                        raise DocumentError(f'Images are limited to {MAX_PAGES} frames/pages.')
                    for i in range(doc.units):
                        im.seek(i)
                        doc.add_image(im.copy())
            elif suffix == '.xlsx':
                _read_xlsx(doc, data)
            elif suffix == '.pptx':
                _read_pptx(doc, data)
            elif suffix == '.odt':
                with zipfile.ZipFile(io.BytesIO(data)) as archive:
                    root = ElementTree.fromstring(archive.read('content.xml'))
                    doc.add_text(' '.join(root.itertext()))
                _archive_images(doc, data, 'Pictures/')
            elif suffix == '.eml':
                message = BytesParser(policy=policy.default).parsebytes(data)
                doc.add_text('Subject: ' + str(message.get('Subject', '')))
                body = message.get_body(preferencelist=('plain', 'html'))
                if body:
                    content = body.get_content()
                    doc.add_text(_html(content) if body.get_content_subtype() == 'html' else content)
                doc.warnings.append('Email body only. Upload attachments separately.')
            else:
                text = _decode(data)
                if suffix in {'.html', '.htm'}:
                    text = _html(text)
                elif suffix == '.rtf':
                    text = rtf_to_text(text)
                elif suffix == '.json':
                    text = json.dumps(json.loads(text), ensure_ascii=False, indent=2)
                elif suffix == '.xml':
                    text = ' '.join(ElementTree.fromstring(text).itertext())
                elif suffix in {'.csv', '.tsv'}:
                    text = '\n'.join(' | '.join(row) for row in csv.reader(io.StringIO(text), delimiter='\t' if suffix == '.tsv' else ','))
                doc.add_text(text)
    except DocumentError:
        raise
    except Exception as exc:
        raise DocumentError('Could not read this file. It may be damaged, password-protected, or saved with the wrong extension. Try exporting it as PDF.') from exc
    doc.text = doc.text.strip()
    if not doc.text and not doc.images:
        raise DocumentError('No readable text or images were found in this file.')
    return doc
